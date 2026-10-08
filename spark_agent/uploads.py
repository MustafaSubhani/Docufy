"""Turn an attached file into markdown notes for a document.

Text and markdown are used as written. Word keeps its headings, lists, and
tables. PDF and slides keep their text in page order. CSV and spreadsheets
become a table. The ingest guards still refuse macros, zip bombs, and long PDFs.
"""

from __future__ import annotations

import csv
import io
import re
import tempfile
from pathlib import Path

from .ingest import IngestError, ingest_path

TEXT = {".md", ".markdown", ".txt"}
SUPPORTED = TEXT | {".pdf", ".docx", ".pptx", ".csv", ".xlsx"}
MAX_TABLE_ROWS = 200
_ANCHOR = re.compile(r"^\[(?:p|s)\d+\]\s*$", re.M)


class UploadReadError(RuntimeError):
    pass


def normalize_newlines(text: str) -> str:
    """LF only. A CRLF string written on Windows otherwise gains a blank line between every row."""
    return text.replace("\r\n", "\n").replace("\r", "\n")


def to_markdown(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise UploadReadError(f"{suffix or name} files are not supported yet")
    if suffix in TEXT:
        return normalize_newlines(data.decode("utf-8", errors="replace").lstrip("\ufeff"))
    if suffix == ".csv":
        return _with_title(name, _csv(data))
    with tempfile.TemporaryDirectory() as raw:
        path = Path(raw) / f"upload{suffix}"
        path.write_bytes(data)
        try:
            doc = ingest_path(path)
        except IngestError as exc:
            raise UploadReadError(str(exc)) from exc
        if suffix == ".docx":
            body = _docx(path)
        elif suffix == ".xlsx":
            body = _sheet(path)
        else:
            body = _pages(doc.text, slides=suffix == ".pptx")
    if not body.strip():
        raise UploadReadError("no text was found in that file")
    return _with_title(name, body)


def _docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    lines: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        level = re.match(r"Heading (\d)", style)
        if style == "Title":
            lines.append(f"# {text}")
        elif level:
            lines.append(f"{'#' * min(int(level.group(1)) + 1, 4)} {text}")
        elif "List Number" in style:
            lines.append(f"1. {text}")
        elif "List" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
        lines.append("")
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(rows) >= 2:
            lines += _table(rows[0], rows[1:]).splitlines() + [""]
    return "\n".join(lines).strip() + "\n"


def _pages(text: str, *, slides: bool) -> str:
    chunks = [part.strip() for part in _ANCHOR.split(text) if part.strip()]
    if not slides:
        paragraphs = []
        for chunk in chunks:
            paragraphs += [" ".join(block.split()) for block in re.split(r"\n\s*\n", chunk) if block.strip()]
        return "\n\n".join(paragraphs) + "\n"
    out = []
    for chunk in chunks:
        lines = [line.strip() for line in chunk.splitlines() if line.strip()]
        if not lines:
            continue
        out.append(f"## {lines[0]}")
        out += [f"- {line}" for line in lines[1:]]
        out.append("")
    return "\n".join(out).strip() + "\n"


def _sheet(path: Path) -> str:
    import zipfile
    from xml.etree import ElementTree

    from .ingest import _shared_strings, _sheet_rows

    with zipfile.ZipFile(path) as package:
        rows = _sheet_rows(ElementTree.fromstring(package.read("xl/worksheets/sheet1.xml")), _shared_strings(package))
    rows = [row for row in rows if any(str(cell).strip() for cell in row)]
    if len(rows) < 2:
        return ""
    return _table(rows[0], rows[1:])


def _csv(data: bytes) -> str:
    rows = list(csv.reader(io.StringIO(data.decode("utf-8", errors="replace").lstrip("\ufeff"))))
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if len(rows) < 2:
        raise UploadReadError("that CSV has no rows under its header")
    return _table(rows[0], rows[1:])


def _table(header: list, rows: list) -> str:
    width = len(header)
    shown = rows[:MAX_TABLE_ROWS]

    def cells(row: list) -> str:
        values = [str(value if value is not None else "").replace("|", "/").strip() for value in list(row)[:width]]
        values += [""] * (width - len(values))
        return "| " + " | ".join(values) + " |"

    lines = [cells(header), "| " + " | ".join("---" for _ in range(width)) + " |"]
    lines += [cells(row) for row in shown]
    if len(rows) > len(shown):
        lines.append("")
        lines.append(f"The first {len(shown)} of {len(rows)} rows are shown.")
    return "\n".join(lines) + "\n"


def _with_title(name: str, body: str) -> str:
    if re.search(r"(?m)^#\s+\S", body):
        return body
    title = re.sub(r"[_\-]+", " ", Path(name).stem).strip() or "Notes"
    return f"# {title[:1].upper() + title[1:]}\n\n{body}"
