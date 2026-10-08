import zipfile

import pymupdf
import pytest
from docx import Document

from spark_agent.ingest import IngestError, ingest_html, ingest_path
from spark_agent.research import build_brief, chunk_document, render_brief


def test_text_is_chunked_and_the_brief_cites_the_anchor(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("Revenue was 12 percent in March 2024.\n\nThe Gulf office led the work.\n", encoding="utf-8")
    doc = ingest_path(path)
    chunks = chunk_document(doc)
    assert chunks and chunks[0].anchor == "p1"
    brief = build_brief([doc], "What was the revenue?")
    assert any("12" in finding.text for finding in brief.findings)
    assert brief.findings[0].anchor == "p1"
    rendered = render_brief(brief)
    assert "waiting for approval" in rendered
    assert "notes:p1" in rendered


def test_csv_profile_keeps_statistics_instead_of_dumping_every_cell(tmp_path):
    path = tmp_path / "rates.csv"
    path.write_text("Item,Qty\nConcrete,10\nPlaster,30\n", encoding="utf-8")
    doc = ingest_path(path)
    table = doc.tables[0]
    assert table.columns == ["Item", "Qty"]
    assert table.row_count == 2
    assert table.numeric["Qty"]["mean"] == 20
    assert "mean 20.00" in doc.text
    assert "Concrete" not in doc.text


def test_pdf_reads_pages_and_flags_hidden_text(tmp_path):
    path = tmp_path / "job.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Visible revenue was 12 percent.", fontsize=12)
    page.insert_text((72, 400), "secret", fontsize=0.2)
    doc.save(path)
    doc.close()
    source = ingest_path(path)
    assert "12 percent" in source.text
    assert source.hidden_text
    assert any("hidden" in warning for warning in source.warnings)


def test_docx_keeps_a_heading_and_a_table(tmp_path):
    path = tmp_path / "scope.docx"
    document = Document()
    document.add_heading("Scope of work", level=1)
    document.add_paragraph("Measure concrete and plaster.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Trade"
    table.rows[0].cells[1].text = "Qty"
    table.rows[1].cells[0].text = "Concrete"
    table.rows[1].cells[1].text = "10"
    document.save(path)
    source = ingest_path(path)
    assert "Scope of work" in source.text
    assert source.tables[0].numeric["Qty"]["mean"] == 10


def test_pptx_and_xlsx_text_come_out_of_the_package(tmp_path):
    slides = tmp_path / "deck.pptx"
    _office_zip(slides, {"ppt/slides/slide1.xml": _slide("Revenue grew in March 2024.")})
    deck = ingest_path(slides)
    assert "Revenue grew" in deck.text
    assert deck.anchors[0].id == "s1"

    book = tmp_path / "book.xlsx"
    _office_zip(
        book,
        {
            "xl/sharedStrings.xml": (
                "<sst><si><t>Item</t></si><si><t>Qty</t></si><si><t>Concrete</t></si></sst>"
            ),
            "xl/worksheets/sheet1.xml": (
                "<worksheet><sheetData>"
                '<row><c t="s"><v>0</v></c><c t="s"><v>1</v></c></row>'
                '<row><c t="s"><v>2</v></c><c><v>10</v></c></row>'
                "</sheetData></worksheet>"
            ),
        },
    )
    sheet = ingest_path(book)
    assert sheet.tables[0].columns == ["Item", "Qty"]
    assert sheet.tables[0].numeric["Qty"]["mean"] == 10


def test_macros_zip_bombs_oversized_files_and_long_pdfs_are_refused(tmp_path):
    macro = tmp_path / "macro.docx"
    _office_zip(macro, {"word/document.xml": "<document/>", "word/vbaProject.bin": "macro"})
    with pytest.raises(IngestError, match="macros"):
        ingest_path(macro)

    bomb = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb, "w", compression=zipfile.ZIP_DEFLATED) as package:
        package.writestr("word/document.xml", b"0" * 400_000)
    with pytest.raises(IngestError, match="expands"):
        ingest_path(bomb)

    big = tmp_path / "big.txt"
    big.write_text("hello", encoding="utf-8")
    with pytest.raises(IngestError, match="limit"):
        ingest_path(big, max_bytes=4)

    long_pdf = tmp_path / "long.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.new_page()
    doc.save(long_pdf)
    doc.close()
    with pytest.raises(IngestError, match="pages"):
        ingest_path(long_pdf, max_pages=1)


def test_html_strips_scripts_and_an_image_is_not_invented(tmp_path):
    page = ingest_html("<html><script>secret()</script><p>Revenue was 12 percent.</p></html>", url="https://example.test")
    assert "Revenue was 12 percent" in page.text
    assert "secret" not in page.text
    image = tmp_path / "shot.png"
    from PIL import Image

    Image.new("RGB", (8, 4), (255, 255, 255)).save(image)
    photo = ingest_path(image)
    assert "8 by 4" in photo.text
    assert any("OCR" in warning for warning in photo.warnings)


def test_a_question_the_files_do_not_answer_is_a_gap(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("The plaster rate is 30 dirhams.\n", encoding="utf-8")
    brief = build_brief([ingest_path(path)], "What is the steel price?")
    assert brief.gaps
    approved = build_brief([ingest_path(path)], "What is the plaster rate?", approved=True)
    assert approved.approved
    assert "approved" in render_brief(approved)


def _office_zip(path, files: dict[str, str]) -> None:
    with zipfile.ZipFile(path, "w") as package:
        for name, text in files.items():
            package.writestr(name, text)


def _slide(text: str) -> str:
    return (
        '<sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f"<a:t>{text}</a:t></sld>"
    )
