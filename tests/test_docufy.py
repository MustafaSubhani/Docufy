import io
import json
import re
import threading
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from spark_agent import ops, panel, secrets, server, store, studio
from spark_agent.draft import clean_section, draft_document
from spark_agent.llm import LLMError, LLMResult, Router, pretty_model
from spark_agent.revise import revise
from spark_agent.uploads import to_markdown

NOTES = (
    "# Site survey\n\n"
    "## Findings\n\nThe north wall needs a second pass.\n\n"
    "## Budget\n\nThe budget is 12,000 AED.\n\n"
    "## Scope\n\nMeasure the yard and the gate.\n"
)


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("SPARK_STUDIO_DIR", str(tmp_path / "studio"))
    monkeypatch.setattr(store, "DB_PATH", tmp_path / "spark.db")
    monkeypatch.setattr(store, "_conn", None)
    secrets.use(secrets.DictVault())
    panel._ollama_probe = lambda _url: {"ok": True, "detail": "reachable"}
    studio._probe_cache.clear()
    yield tmp_path
    _idle()
    secrets.use(None)
    panel._ollama_probe = None
    studio._probe_cache.clear()


def _idle(seconds: float = 20) -> None:
    end = time.time() + seconds
    while studio._jobs.busy() and time.time() < end:
        time.sleep(0.05)


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Site survey", 0)
    doc.add_heading("Findings", 1)
    doc.add_paragraph("The north wall needs a second pass.")
    doc.add_paragraph("Check the gate hinge", style="List Bullet")
    doc.add_heading("Plan", 1)
    doc.add_paragraph("Measure the yard on Monday.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Metres"
    table.rows[1].cells[0].text = "Wall"
    table.rows[1].cells[1].text = "42"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _pdf_bytes() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "The yard visit found a cracked kerb.", fontsize=12)
    data = doc.tobytes()
    doc.close()
    return data


def test_attached_files_become_notes_with_their_structure():
    word = to_markdown("survey.docx", _docx_bytes())
    assert "# Site survey" in word
    assert "## Findings" in word and "- Check the gate hinge" in word
    assert "| Item | Metres |" in word and "| Wall | 42 |" in word
    pdf = to_markdown("yard_visit.pdf", _pdf_bytes())
    assert pdf.startswith("# Yard visit")
    assert "cracked kerb" in pdf and "[p1]" not in pdf
    sheet = to_markdown("costs.csv", b"Item,Cost\nWall,1200\nGate,300\n")
    assert "| Item | Cost |" in sheet and "| Gate | 300 |" in sheet
    guide = to_markdown("guide.md", b"| Export | What you get |\r\n|---|---|\r\n| CSV | Columns |\r\n")
    assert "\r" not in guide
    assert "| Export | What you get |\n|---|---|\n| CSV | Columns |" in guide


def test_a_word_upload_becomes_a_document_named_after_its_title(isolated):
    chat = studio.create_chat()
    detail = studio.send_message(chat["id"], "", upload=("survey.docx", _docx_bytes()), wait=True)
    assert detail["messages"][-1]["body"] == "The document is ready."
    assert detail["messages"][0]["attachment"] == "survey.docx"
    doc = detail["documents"][-1]
    path, name, media = studio.download_path(doc["id"], "pdf")
    assert name == "Site survey.pdf" and media == "application/pdf" and path.is_file()
    _md, md_name, _ = studio.download_path(doc["id"], "md")
    assert md_name == "Site survey.md"


def test_the_open_document_changes_into_new_versions_without_a_model(isolated):
    chat = studio.create_chat()
    first = studio.send_message(chat["id"], NOTES, spark_content=False, wait=True)
    v1 = first["documents"][-1]
    swapped = studio.send_message(chat["id"], "replace north wall with east wall", document_id=v1["id"], wait=True)
    assert swapped["messages"][-1]["body"] == "Replaced north wall with east wall in 1 place. Version 2 is ready."
    v2 = swapped["documents"][-1]
    assert v2["version"] == 2 and v2["parent_id"] == v1["id"] and v2["root_id"] == v1["id"]
    assert "east wall" in studio.download_path(v2["id"], "md")[0].read_text(encoding="utf-8")
    assert "north wall" in studio.download_path(v1["id"], "md")[0].read_text(encoding="utf-8")
    assert [item["id"] for item in studio.list_documents()] == [v2["id"]]
    assert [item["version"] for item in studio.versions(v1["id"])] == [1, 2]

    removed = studio.send_message(chat["id"], "remove the Budget section", document_id=v2["id"], wait=True)
    v3 = removed["documents"][-1]
    text = studio.download_path(v3["id"], "md")[0].read_text(encoding="utf-8")
    assert "12,000" not in text and "Measure the yard" in text
    assert studio.download_path(v3["id"], "pdf")[1] == "Site survey v3.pdf"

    retitled = studio.send_message(chat["id"], "set the title to Yard report", document_id=v3["id"], wait=True)
    assert retitled["documents"][-1]["title"] == "Yard report"
    client = studio.send_message(chat["id"], "make it client facing", document_id=retitled["documents"][-1]["id"], wait=True)
    assert client["documents"][-1]["spark_content"] is True

    count = len(client["documents"])
    asked = studio.send_message(chat["id"], "make the Scope section shorter", document_id=client["documents"][-1]["id"], wait=True)
    assert "Without a model" in asked["messages"][-1]["body"]
    assert len(asked["documents"]) == count
    shorter = studio.send_message(chat["id"], "make it shorter", document_id=client["documents"][-1]["id"], wait=True)
    assert shorter["documents"][-1]["title"] == "Yard report"


def test_a_model_rewrite_changes_only_the_named_section():
    markdown = "---\ntitle: Survey\ninclude_spark_content: false\n---\n\n# 01 Findings\n\n## Findings\n\nThe north wall needs a second pass and a new coat.\n\n# 02 Scope\n\n## Scope\n\nMeasure the yard.\n"
    router = Router(profile="dev", fake=lambda model, messages, schema: "The north wall needs a second pass.")
    result = revise(markdown, "shorten the section Findings", router)
    assert result.summary == "Shortened Findings."
    assert "include_spark_content: false" in result.markdown
    assert "The north wall needs a second pass.\n" in result.markdown and "new coat" not in result.markdown
    assert "Measure the yard." in result.markdown
    unclear = revise(markdown, "make it sound better", router)
    assert unclear.markdown is None and "Findings, Scope" in unclear.summary
    from spark_docs.ir import load_markdown

    findings_id = load_markdown(markdown).parts[0].sections[0].id
    forced = revise(markdown, "make it sound better", router, section_ids=[findings_id])
    assert forced.summary == "Rewrote Findings."
    assert "new coat" not in forced.markdown
    assert "Measure the yard." in forced.markdown


def test_document_sections_and_preview_scoped_revise(isolated, monkeypatch):
    chat = studio.create_chat()
    detail = studio.send_message(chat["id"], NOTES, wait=True)
    doc_id = detail["documents"][-1]["id"]
    listed = studio.document_sections(doc_id)
    assert listed and any(item["label"] == "Findings" for item in listed)
    findings_id = next(item["id"] for item in listed if item["label"] == "Findings")

    def fake(model, messages, schema):
        return "The north wall needs one coat."

    real = studio._router

    def routed(settings):
        router = real(settings)
        router.fake = fake
        return router

    monkeypatch.setattr(studio, "_router", routed)
    revised = studio.send_message(
        chat["id"],
        "shorten",
        document_id=doc_id,
        section_ids=[findings_id],
        model="dev",
        wait=True,
    )
    assert revised["documents"][-1]["id"] != doc_id
    md_path = studio.get_document(revised["documents"][-1]["id"])["md_path"]
    text = Path(md_path).read_text(encoding="utf-8")
    assert "one coat" in text
    assert "Measure the yard" in text


def test_each_message_picks_its_model_and_replies_see_the_chat(isolated, monkeypatch):
    seen: list[str] = []

    def fake(model, messages, schema):
        seen.append(messages[-1]["content"])
        return "Your client is ACME."

    real = studio._router

    def routed(settings):
        router = real(settings)
        router.fake = fake
        return router

    monkeypatch.setattr(studio, "_router", routed)
    chat = studio.create_chat()
    studio.send_message(chat["id"], "The client is ACME Fleet.", model="dev", wait=True)
    detail = studio.send_message(chat["id"], "Who is the client?", model="dev", wait=True)
    reply = detail["messages"][-1]
    assert reply["body"] == "Your client is ACME."
    assert reply["model"] == "dev" and reply["model_label"] == "Gemma4 31B Cloud"
    assert "[[history]]" in seen[-1] and "ACME Fleet" in seen[-1] and "Who is the client?" in seen[-1]
    plain = studio.send_message(chat["id"], "Who is the client?", model="off", wait=True)
    assert plain["messages"][-1]["model_label"] == "Format only"
    with pytest.raises(studio.StudioError, match="choose"):
        studio.send_message(chat["id"], "hi", model="gemini")


def test_a_running_document_can_be_stopped_and_tried_again(isolated, monkeypatch):
    gate = threading.Event()

    def slow(model, messages, schema):
        gate.wait(10)
        raise LLMError("the model is turned off")

    real = studio._router

    def routed(settings):
        router = real(settings)
        router.fake = slow
        return router

    monkeypatch.setattr(studio, "_router", routed)
    chat = studio.create_chat()
    sent = studio.send_message(chat["id"], NOTES, model="dev")
    message = sent["messages"][-1]
    assert message["status"] == "writing"
    stopped = studio.stop_message(chat["id"], message["id"])
    assert stopped["messages"][-1]["status"] == "stopped"
    gate.set()
    _idle()
    after = studio.get_chat(chat["id"])
    assert after["messages"][-1]["body"] == "Stopped." and after["documents"] == []

    studio.retry_message(chat["id"], message["id"], "off")
    _idle()
    done = studio.get_chat(chat["id"])
    assert done["messages"][-1]["status"] == "ready" and len(done["documents"]) == 1
    with pytest.raises(studio.StudioError, match="only a stopped or failed"):
        studio.retry_message(chat["id"], message["id"])


def test_unfinished_replies_rename_search_and_old_logs(isolated):
    chat = studio.create_chat()
    studio.send_message(chat["id"], "Remember the north wall.", wait=True)
    store.conn().execute("UPDATE messages SET status='writing' WHERE role='assistant'")
    store.conn().commit()
    assert studio.recover_messages() == 1
    assert "restarted" in studio.get_chat(chat["id"])["messages"][-1]["body"]
    assert studio.rename_chat(chat["id"], "  Yard   notes ")["title"] == "Yard notes"
    assert [item["id"] for item in studio.list_chats("north wall")] == [chat["id"]]
    assert studio.list_chats("south gate") == []
    store.ensure_panel()
    store.conn().execute("INSERT INTO activity_log (at, kind, detail) VALUES (?, 'note', 'old line')", (time.time() - 40 * 86400,))
    store.conn().commit()
    assert studio.purge_old(30) >= 1
    assert all(row["detail"] != "old line" for row in panel.activity(50))


def test_the_model_menu_says_what_can_run(isolated):
    panel._ollama_probe = lambda _url: {"ok": False, "detail": "not reachable"}
    options = {item["id"]: item for item in studio.model_options()["options"]}
    assert options["off"]["available"] is True
    assert options["dev"]["available"] is False and "Start Ollama" in options["dev"]["reason"]
    assert options["claude"]["reason"] == "Add an Anthropic key in Settings"
    assert options["claude"]["detail"] == "Claude Sonnet 4.5"
    secrets.put("anthropic", "sk-ant-test")
    options = {item["id"]: item for item in studio.model_options()["options"]}
    assert options["claude"]["available"] is True
    studio.save_settings({"local_only": True})
    options = {item["id"]: item for item in studio.model_options()["options"]}
    assert options["claude"]["available"] is False and options["claude"]["reason"] == "Local only is on in Settings"
    assert studio._router({**studio.load_settings(), "profile": "claude", "use_model": True}).profile == "dev"


def test_cloud_calls_record_cost_and_a_cut_off_reply_keeps_its_start(isolated):
    router = Router(profile="claude", fake=lambda model, messages, schema: "Ready.")
    router.complete("writer", "Be brief.", "hello " * 200)
    cost = store.conn().execute("SELECT cost_usd FROM llm_calls ORDER BY id DESC LIMIT 1").fetchone()[0]
    assert cost > 0

    replies = iter(
        [
            LLMResult(text="The fleet is ready. The yard", finish_reason="length"),
            LLMResult(text="The yard is clear.", finish_reason="stop"),
        ]
    )
    cut = Router(profile="dev")
    cut._call = lambda model, messages, schema, max_tokens, temperature: next(replies)
    text = cut.complete("writer", "Be brief.", "Status?").text
    assert text == "The fleet is ready. The yard is clear."
    assert pretty_model("openai/gpt-5-mini") == "GPT 5 Mini"


def test_the_new_routes(isolated, monkeypatch):
    client = TestClient(server.app)
    chat = client.post("/api/chats").json()
    sent = client.post(f"/api/chats/{chat['id']}/messages", json={"text": NOTES, "spark_content": False, "model": "off"})
    assert sent.status_code == 200
    _idle()
    detail = client.get(f"/api/chats/{chat['id']}").json()
    doc = detail["documents"][-1]
    assert client.patch(f"/api/chats/{chat['id']}", json={"title": "Yard"}).json()["title"] == "Yard"
    markdown = client.get(f"/api/documents/{doc['id']}/download?kind=md")
    assert markdown.status_code == 200 and "Site%20survey.md" in markdown.headers["content-disposition"]
    info = client.get(f"/api/documents/{doc['id']}").json()
    assert info["versions"][0]["version"] == 1 and "md_path" not in info
    assert client.get("/api/models").json()["current"] == "off"
    assert client.post(f"/api/chats/{chat['id']}/messages/{'f' * 12}/stop").status_code == 404
    exported = client.get(f"/api/chats/{chat['id']}/export")
    assert exported.status_code == 200 and zipfile.ZipFile(io.BytesIO(exported.content)).namelist()
    imported = client.post("/api/import", files={"file": ("chat.zip", exported.content, "application/zip")})
    assert imported.status_code == 200 and imported.json()["documents"][0]["title"] == "Site survey"
    made = client.post("/api/backups").json()
    assert client.get("/api/backups").json()["backups"][0]["name"] == made["name"]
    monkeypatch.setattr(ops, "ready_checks", lambda: [{"name": "Python", "ok": True, "required": True, "detail": "3.14", "fix": ""}])
    assert client.get("/api/ready").json()["ok"] is True
    assert json.loads(client.get("/api/settings").text)["has_password"] is False

BRIEF = "Write a client proposal for ACME Fleet about fuel tracking on 340 vans"


def _writer(fail: str = ""):
    def fake(model, messages, schema):
        user = messages[-1]["content"]
        if "Plan the document" in user:
            return json.dumps(
                {
                    "archetype": "proposal",
                    "title": "Fuel Tracking for ACME Fleet",
                    "sections": [
                        {"heading": "Executive Summary", "points": ["why it matters"]},
                        {"heading": "Who We Are", "points": []},
                        {"heading": "Scope of Work", "points": ["units per van"]},
                        {"heading": "Commercial Proposal", "points": ["price to be confirmed"]},
                    ],
                }
            )
        found = re.search(r'Write the section "([^"]+)"', user)
        if found:
            if found.group(1) == fail:
                raise LLMError("the writer timed out")
            return f"## {found.group(1)}\n\n{found.group(1)} covers the 340 vans \u2014 fuel cards and GPS."
        raise LLMError("layout is left to the notes")

    return fake


def _routed(monkeypatch, fake):
    real = studio._router

    def routed(settings):
        router = real(settings)
        router.fake = fake
        return router

    monkeypatch.setattr(studio, "_router", routed)


class Recorder:
    def __init__(self):
        self.rows = {}

    def start(self, key, label, kind="step", parent="", detail=""):
        self.rows[key] = {"label": label, "kind": kind, "parent": parent, "status": "running"}

    def finish(self, key, detail="", status="done"):
        self.rows[key]["status"] = status
        self.rows[key]["detail"] = detail

    def check(self):
        pass


def test_a_brief_is_written_into_a_new_document_by_section_agents(isolated, monkeypatch):
    _routed(monkeypatch, _writer())
    chat = studio.create_chat()
    detail = studio.send_message(chat["id"], BRIEF, model="dev", spark_content=True, wait=True)
    reply = detail["messages"][-1]
    assert reply["status"] == "ready" and reply["body"].startswith("The draft is ready.")
    doc = detail["documents"][-1]
    assert doc["title"] == "Fuel Tracking for ACME Fleet"
    notes = (studio.studio_dir() / "docs" / doc["id"] / "notes.md").read_text(encoding="utf-8")
    assert "## Scope of Work" in notes and "340 vans, fuel cards and GPS." in notes
    assert "## Who We Are" not in notes and "include_spark_content: true" in notes and "doc_type: PROPOSAL" in notes
    steps = {step["key"]: step for step in reply["activity"]}
    assert [steps[key]["label"] for key in ("outline", "write", "assemble")] == ["Outlining the document", "Writing the sections", "Assembling the draft"]
    agents = [step for step in reply["activity"] if step["kind"] == "agent" and step["parent"] == "write"]
    assert [step["label"] for step in agents] == ["Executive Summary", "Scope of Work", "Commercial Proposal"]
    assert all(step["status"] == "done" and step["ended"] >= step["started"] for step in reply["activity"])
    assert {"read", "plan", "layout", "render", "preview"} <= set(steps)


def test_a_brief_without_a_model_waits_for_one_and_can_be_tried_again(isolated, monkeypatch):
    chat = studio.create_chat()
    blocked = studio.send_message(chat["id"], "Draft a scope of work for the yard survey", model="off", wait=True)
    message = blocked["messages"][-1]
    assert message["status"] == "blocked" and message["body"].startswith("Writing a new document needs a model.")
    _routed(monkeypatch, _writer())
    studio.retry_message(chat["id"], message["id"], "dev")
    _idle()
    done = studio.get_chat(chat["id"])
    assert done["messages"][-1]["status"] == "ready" and len(done["documents"]) == 1
    assert done["messages"][-1]["model"] == "dev"


def test_cloud_agents_write_in_parallel_and_a_failed_section_is_marked():
    router = Router(profile="claude", fake=_writer(fail="Scope of Work"))
    events = Recorder()
    draft = draft_document(router, BRIEF, spark=False, events=events)
    assert draft.missing == ["Scope of Work"]
    assert "This section is to be confirmed." in draft.markdown
    assert events.rows["write-1"]["status"] == "failed" and events.rows["write-0"]["status"] == "done"
    assert events.rows["write"]["detail"] == "2 of 3 sections written"
    assert "include_spark_content: false" in draft.markdown
    assert clean_section("```markdown\n# Scope\n\nRange 3\u20135 \u2014 fine.\n```", "Scope") == "Range 3 to 5, fine."


def test_a_web_question_searches_and_a_following_draft_uses_the_pages(isolated, monkeypatch):
    from spark_agent.deep import Leaf, MemoryConnector, Page, SearchHit, run_research

    def searched(goal):
        hit = SearchHit(url="https://example.gov/fauji", title="Fauji Foundation", connector="ddg")

        def fetch(url):
            return Page(url=url, title="Fauji Foundation", text="Fauji Foundation employs 80000 people in 2024.")

        return run_research(
            goal,
            depth="quick",
            connectors=[MemoryConnector("ddg", [hit])],
            fetch=fetch,
            year=2026,
            tree=[Leaf(id="q1", question=goal, evidence="official figures", route="web")],
        )

    monkeypatch.setattr(studio, "_run_web", searched)
    chat = studio.create_chat()
    asked = studio.send_message(
        chat["id"],
        "identify a problem at Fauji Foundation with deep research from the internet",
        model="off",
        wait=True,
    )
    reply = asked["messages"][-1]
    assert reply["status"] == "ready" and asked["documents"] == []
    assert "80000" in reply["body"] and "https://example.gov/fauji" in reply["body"]
    assert any(step["label"] == "Searching the web" for step in reply["activity"])

    seen = []

    def fake(model, messages, schema):
        seen.append(messages[-1]["content"])
        return _writer()(model, messages, schema)

    _routed(monkeypatch, fake)
    drafted = studio.send_message(chat["id"], "help me make a spark ai document about progress tracking", model="dev", wait=True)
    assert drafted["documents"]
    assert drafted["messages"][-1]["body"].startswith("I searched the web, then wrote the draft.")
    assert any("https://example.gov/fauji" in text for text in seen)


def test_the_web_switch_searches_a_plain_question(isolated, monkeypatch):
    from spark_agent.deep import Leaf, MemoryConnector, Page, SearchHit, run_research

    def searched(goal):
        hit = SearchHit(url="https://example.gov/pakistan", title="Pakistan industry", connector="ddg")

        def fetch(url):
            return Page(url=url, title="Pakistan industry", text="Large firms in Pakistan employed 2400000 people in 2023.")

        return run_research(
            goal,
            depth="quick",
            connectors=[MemoryConnector("ddg", [hit])],
            fetch=fetch,
            year=2026,
            tree=[Leaf(id="q1", question=goal, evidence="official figures", route="web")],
        )

    monkeypatch.setattr(studio, "_run_web", searched)
    chat = studio.create_chat()
    detail = studio.send_message(chat["id"], "help me research some areas in Pakistan", model="off", web=True, wait=True)
    assert detail["messages"][0]["web"] is True
    assert "2400000" in detail["messages"][-1]["body"]
    assert any(step["label"] == "Searching the web" for step in detail["messages"][-1]["activity"])


def test_a_chat_search_reply_names_the_query(isolated):
    from spark_agent.deep import ChatResult, Page

    run = ChatResult(
        query="AI solutions Pakistan organizations",
        pages=[Page(url="https://example.com/pk", title="AI in Pakistan", text="Large companies in Pakistan are starting to use AI.")],
        findings=["Large companies in Pakistan are starting to use AI for credit checks."],
    )
    body = studio._chat_findings(run)
    assert body.startswith("I searched for AI solutions Pakistan organizations.")
    assert "[AI in Pakistan](https://example.com/pk)" in body and "Closest pages" in body
    assert "Findings" not in body and "doi.org" not in body


def test_the_model_chooses_the_search_and_writes_the_answer(isolated, monkeypatch):
    from spark_agent.deep import ChatResult, Page

    seen = []

    def searched(goal, query="", **kwargs):
        seen.append(query or goal)
        return ChatResult(
            query=query or goal,
            pages=[Page(url="https://example.com/engro", title="Engro and Baidu launch AI in Pakistan", text="")],
            findings=[""],
        )

    def fake(_model, messages, schema):
        user = messages[-1]["content"]
        if schema:
            assert "pakistan" in user.lower()
            return json.dumps({"queries": ["AI adoption Pakistan banking", "Engro AI Pakistan"]})
        assert "https://example.com/engro" in user
        assert "Excerpt:" in user or "Engro" in user
        return (
            "---openui\n"
            'root = ResearchView("AI solutions in Pakistan", "Banking and Engro are named in the pages.", '
            '[{problem: "Adoption", aiUse: "Analytics", caveat: "Data quality"}], '
            '"More telecom detail was missing.", '
            '[{title: "Engro and Baidu launch AI in Pakistan", url: "https://example.com/engro"}])\n'
            "---end"
        )

    monkeypatch.setattr(studio, "_run_web", searched)
    _routed(monkeypatch, fake)
    chat = studio.create_chat()
    detail = studio.send_message(chat["id"], "help me research ai solutions in pakistan", model="dev", web=True, wait=True)
    body = detail["messages"][-1]["body"]
    assert "---openui" in body and "ResearchView" in body
    assert "https://example.com/engro" in body
    assert "AI adoption Pakistan banking" in seen
    labels = [step["label"] for step in detail["messages"][-1]["activity"]]
    assert "Quick lookup" in labels and "Research" in labels


def test_deep_research_uses_the_verified_pipeline(isolated, monkeypatch):
    from spark_agent.deep import DeepRun, Page
    from spark_agent.research import Finding, ResearchBrief

    def fake_run(goal, **kwargs):
        page = Page(
            url="https://example.com/energy",
            title="Pakistan energy AI",
            text="Predictive maintenance and leak detection are common AI uses in energy. " * 8,
        )
        brief = ResearchBrief(
            question=goal,
            findings=[Finding(text="Predictive maintenance reduces downtime.", anchor="", source=page.url)],
        )
        return DeepRun(goal=goal, depth="standard", pages=[page], brief=brief, checked_on="2026-01-01")

    def fake(_model, messages, schema):
        if schema:
            return json.dumps({"queries": ["AI Pakistan energy"]})
        assert "Confirmed findings" in messages[-1]["content"]
        return (
            "---openui\n"
            'root = ResearchView("Energy AI in Pakistan", "Predictive maintenance appears in the notes.", '
            '[{problem: "Downtime", aiUse: "Predictive maintenance", caveat: "Sensor coverage"}], '
            '"Broader policy context was thin.", '
            '[{title: "Pakistan energy AI", url: "https://example.com/energy"}])\n'
            "---end"
        )

    monkeypatch.setattr("spark_agent.deep.run_research", fake_run)
    _routed(monkeypatch, fake)
    chat = studio.create_chat()
    detail = studio.send_message(
        chat["id"],
        "do a deep research on ai solutions for big organisations in pakistan",
        model="dev",
        wait=True,
    )
    body = detail["messages"][-1]["body"]
    labels = [step["label"] for step in detail["messages"][-1]["activity"]]
    assert "ResearchView" in body
    assert "https://example.com/energy" in body
    assert "Deep research" in labels


def test_research_openui_prompt_is_shipped():
    path = studio.ROOT / "spark_agent" / "prompts" / "research_openui_prompt.txt"
    text = path.read_text(encoding="utf-8")
    assert "ResearchView" in text and "root =" in text


def test_normalize_openui_fenced_and_named_args():
    raw = (
        "```openui\n"
        "root = ResearchView(\n"
        'ask: "What industries?",\n'
        'context: "Two hubs.",\n'
        'insights: [{ problem: "A", aiUse: "B" }],\n'
        'gaps: "More detail.",\n'
        'sources: [{ title: "T", url: "https://example.com" }]\n'
        ")\n"
        "```"
    )
    out = studio._normalize_openui_lang(raw)
    assert out.startswith("---openui")
    assert "ask:" not in out.split("---end")[0]
    assert "ResearchView(" in out and '"What industries?"' in out


def test_normalize_openui_object_equals_to_colons():
    raw = (
        "---openui\n"
        'root = ResearchView("a", "b", [{ problem = "p", aiUse = "u" }], "g", '
        '[{ title = "t", url = "https://example.com" }])\n'
        "---end"
    )
    out = studio._normalize_openui_lang(raw)
    assert "problem:" in out and "problem =" not in out
    assert "aiUse:" in out and "title:" in out


def test_research_answer_strips_brands_not_in_sources():
    from spark_agent.deep import Page

    pages = [Page(url="https://example.com/oil", title="Oil AI", text="upstream maintenance and safety")]
    raw = "Line about [Genspark AI Docs](https://genspark.ai).\nLine about [Oil](https://example.com/oil)."
    out = studio._guard_research_answer(raw, pages)
    assert "genspark" not in out.lower()
    assert "example.com/oil" in out


def test_a_quick_question_stays_a_lookup_and_a_dig_goes_deep():
    from spark_agent.scout import research_depth

    assert research_depth("help me research ai solutions in pakistan") == "quick"
    assert research_depth("do some deep research from the internet") == "deep"
    assert research_depth("dig into the cement plants") == "deep"


def test_a_reply_does_not_repeat_the_task_marker(isolated, monkeypatch):
    _routed(monkeypatch, lambda model, messages, schema: "[[task]]\nThe yard is fine.")
    chat = studio.create_chat()
    detail = studio.send_message(chat["id"], "What is the status of the yard?", model="dev", wait=True)
    assert detail["messages"][-1]["body"] == "The yard is fine."


def test_briefs_notes_and_questions_are_told_apart():
    assert studio.is_brief("Write a proposal for ACME")
    assert studio.is_brief("Can you draft a report on the yard visit?")
    assert not studio.is_brief("# Notes\n\n## Findings\n\nThe wall.")
    assert not studio.is_brief("Hello there")
    assert not studio.is_brief("Write a proposal", "notes.docx")
    assert studio._kind("write a new proposal for Beta Co", "", "a" * 12) == "draft"
    assert studio._kind("make it shorter", "", "a" * 12) == "revise"
    assert studio.wants_web("do some deep research from the internet")
    assert studio.wants_web("find it from the internet research")
    assert studio.wants_web("i want you to resesearch on internet")
    assert not studio.wants_web("help me research some areas in Pakistan")
    assert not studio.wants_web("Write a proposal for ACME")
    assert studio.research_intent("help me research some areas in Pakistan")
    assert studio.research_intent("help me reearch about oil and gas and common problems we can solve using ai")
    assert studio._kind("do some deep research from the internet", "", "") == "research"
    assert studio._kind("help me research some areas in Pakistan", "", "") == "research"
    assert studio._kind("help me research some areas in Pakistan", "", "", web=True) == "research"
