"""Spark-branded DOCX export (editable companion to the PDF)."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor

from .images import raster_path_for_pdf
from .layout import fit_image
from .parser import format_kicker
from .theme import LOGO_PATH, Meta

NAVY = RGBColor(0x1A, 0x1A, 0x2E)
PURPLE = RGBColor(0x6B, 0x4E, 0xFF)
BODY = RGBColor(0x3B, 0x3B, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_DEEP = "0D1B3E"
CALLOUT = "E5EAF1"
ALT = "F5F7FA"


def build_docx(meta: Meta, blocks: list[dict], output: Path) -> Path:
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.6)
    section.different_first_page_header_footer = True
    if meta.top_bar:
        _page_top_border(section)

    _set_run_font(doc.styles["Normal"], "Arial", 10, BODY)
    for style_name, size, color, bold in (
        ("Heading 1", 9, PURPLE, True),
        ("Heading 2", 22, NAVY, True),
        ("Heading 3", 11, NAVY, True),
    ):
        try:
            _set_run_font(doc.styles[style_name], "Arial", size, color, bold=bold)
        except KeyError:
            pass

    if meta.inner_logo:
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if LOGO_PATH.exists():
            run = hp.add_run()
            run.add_picture(str(LOGO_PATH), width=Inches(1.25))

    if meta.inner_footer:
        footer = section.footer
        footer.is_linked_to_previous = False
        _footer_bar(footer, meta.footer)

        first_footer = section.first_page_footer
        first_footer.is_linked_to_previous = False
        fp0 = first_footer.paragraphs[0]
        fp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = fp0.add_run(meta.footer)
        _font(r0, 8, RGBColor(0x88, 0x88, 0x88))

    if meta.has_cover():
        _cover(doc, meta)

    for block in blocks:
        _emit(doc, block, meta)

    doc.save(str(output))
    return output


def _cover(doc: Document, meta: Meta) -> None:
    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        run = p.add_run()
        run.add_picture(str(LOGO_PATH), width=Inches(1.5))
        p.paragraph_format.space_after = Pt(18)
    if meta.doc_type:
        p = doc.add_paragraph()
        run = p.add_run(meta.doc_type.upper())
        _font(run, 9, PURPLE, bold=True)
        p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph()
    run = p.add_run(meta.title)
    _font(run, 32, NAVY, bold=True)
    p.paragraph_format.space_after = Pt(6)
    if meta.subtitle:
        p = doc.add_paragraph()
        run = p.add_run(meta.subtitle)
        _font(run, 16, RGBColor(0x44, 0x44, 0x44))
        p.paragraph_format.space_after = Pt(10)
    _hr(doc, "6C3FE0", size="12")
    for label, value in (
        ("Prepared for", meta.prepared_for),
        ("Prepared by", meta.prepared_by),
        ("Date", meta.date),
    ):
        if value:
            p = doc.add_paragraph()
            run = p.add_run(f"{label}:  ")
            _font(run, 10, NAVY, bold=True)
            run = p.add_run(value)
            _font(run, 10, BODY)
            p.paragraph_format.space_after = Pt(2)
    if meta.confidential:
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        _shade(cell, "F8F7FE")
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(meta.confidential_text)
        _font(run, 9, RGBColor(0x64, 0x74, 0x8B), italic=True)
        _set_cell_padding(cell, 120)
    doc.add_page_break()


def _emit(doc: Document, block: dict, meta: Meta) -> None:
    kind = block.get("kind")
    if kind == "heading":
        level = min(block.get("level", 1), 3)
        text = block.get("text") or ""
        if level == 1:
            p = doc.add_paragraph(format_kicker(text), style="Heading 1")
            for run in p.runs:
                run.font.color.rgb = PURPLE
                run.font.size = Pt(8)
                run.bold = True
                run.font.name = "Arial"
            _bottom_border(p, "C8C8C8", "8")
            p.paragraph_format.space_after = Pt(10)
            return
        if level == 2:
            size = 14 if block.get("compact") else (26 if len(text) < 42 else 20)
            p = doc.add_paragraph(text)
            for run in p.runs:
                _font(run, size, NAVY, bold=True)
            p.paragraph_format.space_after = Pt(10)
            return
        if level == 3:
            size = 14 if block.get("module") else 11
            color = PURPLE if block.get("accent") else NAVY
            p = doc.add_paragraph(text)
            for run in p.runs:
                _font(run, size, color, bold=True)
            p.paragraph_format.space_after = Pt(6)
            return
        p = doc.add_paragraph(text, style=f"Heading {level}")
        if block.get("accent"):
            for run in p.runs:
                run.font.color.rgb = PURPLE
        return
    if kind == "para":
        p = doc.add_paragraph()
        _add_inline(p, block.get("text") or "")
        return
    if kind == "list":
        style = "List Number" if block.get("ordered") else "List Bullet"
        for item in block.get("items") or []:
            p = doc.add_paragraph(style=style)
            _add_inline(p, item.get("text") or "")
        return
    if kind == "table":
        headers = block.get("headers") or []
        rows = block.get("rows") or []
        table = doc.add_table(rows=1 + len(rows), cols=max(len(headers), 1))
        table.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = _plain(h)
            _shade(cell, NAVY_DEEP)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = WHITE
                    run.font.size = Pt(9)
                    run.font.name = "Arial"
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = table.rows[r_i + 1].cells[c_i]
                cell.text = _plain(val)
                if r_i % 2 == 1:
                    _shade(cell, ALT)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.name = "Arial"
        doc.add_paragraph()
        return
    if kind == "callout":
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        _shade(cell, CALLOUT)
        cell.text = ""
        _set_cell_padding(cell, 140)
        first = True
        if block.get("title"):
            p = cell.paragraphs[0]
            run = p.add_run(block["title"])
            _font(run, 11, NAVY, bold=True)
            first = False
        for child in block.get("blocks") or []:
            if child.get("kind") == "para":
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                _add_inline(p, child.get("text") or "")
            elif child.get("kind") == "list":
                for item in child.get("items") or []:
                    p = cell.add_paragraph(style="List Bullet")
                    _add_inline(p, item.get("text") or "")
                    first = False
        doc.add_paragraph()
        return
    if kind == "columns":
        cols = block.get("columns") or []
        if not cols:
            return
        table = doc.add_table(rows=1, cols=len(cols))
        for i, col in enumerate(cols):
            cell = table.rows[0].cells[i]
            cell.text = ""
            if col.get("title"):
                p = cell.paragraphs[0]
                run = p.add_run(col["title"])
                _font(run, 11, PURPLE if col.get("accent") else NAVY, bold=True)
            for child in col.get("blocks") or []:
                if child.get("kind") == "para":
                    p = cell.add_paragraph()
                    _add_inline(p, child.get("text") or "")
        doc.add_paragraph()
        return
    if kind == "cards":
        cards = block.get("cards") or []
        chunk = 4 if len(cards) == 4 else 2
        for i in range(0, len(cards), chunk):
            row_cards = cards[i : i + chunk]
            table = doc.add_table(rows=1, cols=len(row_cards))
            for c_i, card in enumerate(row_cards):
                cell = table.rows[0].cells[c_i]
                _shade(cell, CALLOUT)
                cell.text = ""
                p = cell.paragraphs[0]
                _add_card_title(p, card.get("title") or "")
                for child in card.get("blocks") or []:
                    if child.get("kind") == "para":
                        p = cell.add_paragraph()
                        _add_inline(p, child.get("text") or "", size=9.5)
            doc.add_paragraph()
        return
    if kind == "stageflow":
        steps = block.get("steps") or []
        fills = {
            "bot": "D6E8FF",
            "ai": "E8E4FF",
            "output": "DFF0E8",
            "gov": "F0E8DC",
            "process": "E8E4FF",
        }
        for step in steps:
            table = doc.add_table(rows=1, cols=2)
            layer = step.get("layer") or "process"
            left = table.rows[0].cells[0]
            right = table.rows[0].cells[1]
            _shade(left, fills.get(layer, fills["process"]))
            _shade(right, "F4F6F9")
            left.text = ""
            right.text = ""
            p = left.paragraphs[0]
            run = p.add_run(step.get("title") or "")
            _font(run, 9.5, NAVY, bold=True)
            if step.get("detail"):
                p = left.add_paragraph()
                _add_inline(p, step["detail"], size=8.5)
            p = right.paragraphs[0]
            run = p.add_run(step.get("cap_title") or "")
            _font(run, 9, NAVY, bold=True)
            if step.get("cap_detail"):
                p = right.add_paragraph()
                _add_inline(p, step["cap_detail"], size=8.5)
            doc.add_paragraph()
        return
    if kind == "hr":
        _hr(doc, "C8C8C8", "8")
        return
    if kind == "image":
        path = Path(block.get("path") or "")
        if path.is_file():
            embed = Path(raster_path_for_pdf(path))
            w_pt, h_pt = _image_size(embed, block)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run()
            run.add_picture(str(embed), width=Inches(w_pt / 72), height=Inches(h_pt / 72))
        else:
            p = doc.add_paragraph()
            run = p.add_run(f"Image not found: {block.get('alt') or path.name}")
            _font(run, 9, BODY, italic=True)
        return
    if kind == "code":
        p = doc.add_paragraph(block.get("text") or "")
        for run in p.runs:
            run.font.name = "Courier New"
            run.font.size = Pt(8)
        return
    if kind == "contact":
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        _shade(cell, NAVY_DEEP)
        cell.text = ""
        _set_cell_padding(cell, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("To proceed, please contact us:")
        _font(run, 11, WHITE, bold=True)
        c = meta.contact or {}
        line = "  |  ".join(x for x in (c.get("email"), c.get("phone"), c.get("website")) if x)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        _font(run, 9, WHITE)
        if c.get("address"):
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(c["address"])
            _font(run, 9, WHITE)
        doc.add_paragraph()
        return
    if kind == "pagebreak":
        doc.add_page_break()


def _image_size(path: Path, block: dict) -> tuple[float, float]:
    asset_scale = float(block.get("asset_scale") or 1)
    is_diagram = bool(block.get("diagram"))
    wide = bool(block.get("diagram_wide"))
    try:
        from PIL import Image as PILImage

        with PILImage.open(path) as img:
            return fit_image(
                img.width,
                img.height,
                asset_scale,
                diagram=is_diagram,
                diagram_wide=wide,
            )
    except Exception:
        return fit_image(1200, 800, asset_scale, diagram=is_diagram, diagram_wide=wide)


CARD_NUM_RE = re.compile(r"^(\d{2})\s+(.+)$")


def _add_card_title(p, title: str) -> None:
    m = CARD_NUM_RE.match(title.strip())
    if m:
        run = p.add_run(m.group(1))
        _font(run, 10, PURPLE, bold=True)
        run = p.add_run(f"  {m.group(2)}")
        _font(run, 10, NAVY, bold=True)
        return
    run = p.add_run(title)
    _font(run, 10, NAVY, bold=True)


INLINE_TOKEN_RE = re.compile(
    r"(\*\*.+?\*\*|\*(?!\*).+?\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s\])<>\"']+)"
)
MD_LINK_RE = re.compile(r"^\[([^\]]+)\]\(([^)]+)\)$")


def _add_hyperlink(paragraph, text: str, url: str, size: float = 10) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "6B4EFF")
    r_pr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(r_fonts)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _add_inline(p, text: str, size: float = 10) -> None:
    if not text:
        _font(p.add_run(""), size, BODY)
        return
    pos = 0
    for m in INLINE_TOKEN_RE.finditer(text):
        if m.start() > pos:
            run = p.add_run(text[pos : m.start()])
            _font(run, size, BODY)
        token = m.group(0)
        if token.startswith("**"):
            run = p.add_run(token[2:-2])
            _font(run, size, BODY, bold=True)
        elif token.startswith("`"):
            run = p.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif token.startswith("*"):
            run = p.add_run(token[1:-1])
            _font(run, size, BODY, italic=True)
        elif token.startswith("["):
            lm = MD_LINK_RE.match(token)
            if lm:
                _add_hyperlink(p, lm.group(1), lm.group(2), size=size)
            else:
                run = p.add_run(token)
                _font(run, size, BODY)
        elif token.startswith("http"):
            url = token.rstrip(".,;)")
            _add_hyperlink(p, url, url, size=size)
        else:
            run = p.add_run(token)
            _font(run, size, BODY)
        pos = m.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        _font(run, size, BODY)


def _plain(text: str) -> str:
    import re

    return re.sub(r"[*_`]+", "", text or "")


def _font(run, size, color, bold=False, italic=False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")


def _set_run_font(style, name, size, color, bold=None) -> None:
    font = style.font
    font.name = name
    font.size = Pt(size)
    font.color.rgb = color
    if bold is not None:
        font.bold = bold
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def _hr(doc: Document, color: str, size: str = "8") -> None:
    p = doc.add_paragraph()
    _bottom_border(p, color, size)


def _bottom_border(p, color: str, size: str = "8") -> None:
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _page_top_border(section) -> None:
    sectPr = section._sectPr
    pgBorders = OxmlElement("w:pgBorders")
    pgBorders.set(qn("w:offsetFrom"), "page")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "24")
    top.set(qn("w:space"), "0")
    top.set(qn("w:color"), "6B4EFF")
    pgBorders.append(top)
    sectPr.append(pgBorders)


def _footer_bar(footer, text: str) -> None:
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _bottom_border(p, "C8C8C8", "6")
    run = p.add_run(text)
    _font(run, 8, RGBColor(0x88, 0x88, 0x88))
    run = p.add_run("\t")
    run = p.add_run()
    _add_page_number(run)
    _font(run, 8, RGBColor(0x88, 0x88, 0x88))
    tabs = p.paragraph_format.tab_stops
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm as _Cm

    tabs.add_tab_stop(_Cm(16.5), WD_TAB_ALIGNMENT.RIGHT)


def _add_page_number(run) -> None:
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_end)


def _set_cell_padding(cell, twips: int = 100) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(twips))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)
